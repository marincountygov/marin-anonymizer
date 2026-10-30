#!/usr/bin/env python3
"""Check browser-generated CSV, JSON, and XLSX downloads with invented records only.

The same real application JavaScript and vendored parsers are exercised in both modes:
normal localhost HTTP, or a no-network in-memory page for restricted environments.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
from contextlib import contextmanager, nullcontext
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import shutil
from threading import Thread
from urllib.parse import urlsplit

from html_fixture import remove_from_html

ROOT = Path(__file__).resolve().parents[1]
CSV_FIXTURE = ROOT / "tests/fixtures/synthetic-records.csv"
PREFIX = "/marin-anonymizer-test/"
REDACTED = "---REDACTED---"
SCRIPTS = (
    "vendor/marinos/marinos.js",
    "libs/papaparse.min.js",
    "libs/xlsx.full.min.js",
    "libs/crypto-js.min.js",
    "libs/faker.min.js",
    "assets/app.js",
)


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def do_GET(self) -> None:
        # Test under a URL subpath so root-relative references do not go unnoticed.
        if not self.path.startswith(PREFIX):
            self.send_error(404)
            return
        self.path = "/" + self.path[len(PREFIX):]
        super().do_GET()

    def log_message(self, *_args) -> None:
        pass


@contextmanager
def serve():
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}{PREFIX}"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def fixture_rows() -> list[dict[str, str]]:
    with CSV_FIXTURE.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def workbook_bytes(rows: list[dict[str, str]]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    main = book.active
    main.title = "Current"
    archive = book.create_sheet("Archive")
    archive.sheet_state = "hidden"
    for sheet, records in ((main, rows), (archive, [rows[1], rows[3]])):
        sheet.append(list(rows[0]))
        for record in records:
            sheet.append([record[key] for key in rows[0]])
        for cell in sheet["A"]:
            cell.number_format = "@"  # IDs such as 0001 remain text, never numbers.
    result = io.BytesIO()
    book.save(result)
    return result.getvalue()


def read_workbook(data: bytes) -> tuple[dict[str, list[dict[str, str]]], dict[str, str]]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    try:
        output = {}
        states = {}
        for sheet in book.worksheets:
            values = list(sheet.values)
            headers = list(values[0])
            output[sheet.title] = [
                dict(zip(headers, ("" if value is None else str(value) for value in row)))
                for row in values[1:]
            ]
            states[sheet.title] = sheet.sheet_state
        return output, states
    finally:
        book.close()


def assert_transformation(source: list[dict[str, str]], actual: list[dict[str, str]], label: str) -> None:
    require(len(actual) == len(source), f"{label}: row count changed")
    headers = list(source[0])
    for number, (original, output) in enumerate(zip(source, actual), start=1):
        require(list(output) == headers, f"{label}, row {number}: field names or order changed")
        for field in headers:
            value = original[field]
            if field == "person_name":
                expected = REDACTED if value else ""
            elif field == "email":
                expected = hashlib.sha256(value.encode("utf-8")).hexdigest() if value else ""
            else:
                expected = value
            require(output[field] == expected,
                    f"{label}, row {number}, {field}: expected {expected!r}, got {output[field]!r}")


def make_page(context, base_url: str | None):
    page = context.new_page()
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))

    if base_url is not None:
        page.goto(base_url, wait_until="load")
    else:
        # No browser navigation is possible in some locked-down runners. Inject
        # the actual checked-in HTML, styles, libraries, and app JS into a page.
        html = (ROOT / "index.html").read_text(encoding="utf-8")
        html = remove_from_html(html, "scripts", "stylesheets")
        page.set_content(html)
        for css in ("vendor/marinos/marinos.css", "assets/app.css"):
            page.add_style_tag(content=(ROOT / css).read_text(encoding="utf-8"))
        # The shell may request its optional catalog or updates. Stub those
        # network reads only: no test data should leave the browser.
        page.evaluate("""() => {
          window.__fixtureFetches = [];
          window.fetch = async (input, init) => {
            const method = (init?.method || 'GET').toUpperCase();
            window.__fixtureFetches.push({method, url: String(input)});
            if (method !== 'GET') throw new Error('Unexpected network write');
            return new Response('{}', {status: 200,
              headers: {'Content-Type': 'application/json'}});
          };
        }""")
        for path in SCRIPTS:
            page.add_script_tag(content=(ROOT / path).read_text(encoding="utf-8"))

    page.wait_for_function("() => !!window.Papa && !!window.XLSX && !!window.CryptoJS")
    page.locator("#choose-file-button").wait_for(state="visible")
    require(page.locator("#app-status-message").count() == 1, "Missing app status region")
    return page, errors


def anonymize(page, filename: str, mime: str, content: bytes) -> bytes:
    require(page.locator("#process-button").is_disabled(), "Download starts enabled without a file")
    page.locator("#file-input").set_input_files({
        "name": filename, "mimeType": mime, "buffer": content,
    })
    page.locator('#mapping-rows select[data-column="person_name"]').wait_for()
    if filename.endswith(".xlsx"):
        selector = page.locator("#worksheet-select")
        require(selector.is_visible() and selector.input_value() == "all",
                "XLSX test must include all worksheets, including the hidden sheet")
    require(page.locator("#process-button").is_disabled(), "Download enabled without chosen fields")
    require(page.locator("#mapping-rows select").count() == 5, "Missing or unexpected mapping fields")
    page.locator('#mapping-rows select[data-column="person_name"]').select_option("redact")
    page.locator('#mapping-rows select[data-column="email"]').select_option("hash")
    for field in ("record_id", "service", "notes"):
        require(page.locator(f'#mapping-rows select[data-column="{field}"]').input_value() == "none",
                f"Unselected field {field} unexpectedly has a transform")
    page.locator("#process-button").wait_for(state="visible")
    with page.expect_download(timeout=15000) as download_info:
        page.locator("#process-button").click()
    download = download_info.value
    expected_filename = filename.rsplit(".", 1)[0] + "-anonymized." + filename.rsplit(".", 1)[1]
    require(download.suggested_filename == expected_filename,
            f"Incorrect download filename: {download.suggested_filename!r}")
    path = download.path()
    require(path is not None, "No downloadable file was produced")
    data = path.read_bytes()
    require(bool(data), "Downloaded output was empty")
    page.get_by_text(f"Download started: {expected_filename}", exact=False).wait_for()
    return data


def run(in_memory: bool) -> None:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as error:
        raise RuntimeError("Install test dependencies: python3 -m pip install -r tests/requirements.txt") from error

    rows = fixture_rows()
    require(len(rows) == 4 and len(rows[0]) == 5, "Synthetic test fixture unexpectedly changed")
    local = nullcontext(None) if in_memory else serve()
    with local as base_url, sync_playwright() as playwright:
        executable = (os.environ.get("CHROMIUM_EXECUTABLE") or shutil.which("chromium")
                      or shutil.which("chromium-browser") or shutil.which("google-chrome"))
        options = {"headless": True, "args": ["--no-sandbox", "--disable-dev-shm-usage"]}
        if executable:
            options["executable_path"] = executable
        browser = playwright.chromium.launch(**options)
        try:
            context = browser.new_context(accept_downloads=True)
            context.set_default_timeout(12000)
            network_writes = []
            context.on("request", lambda request: network_writes.append(f"{request.method} {request.url}")
                       if request.method.upper() not in ("GET", "HEAD") else None)
            if base_url:
                origin = urlsplit(base_url).netloc

                def intercept(route) -> None:
                    url = urlsplit(route.request.url)
                    if url.netloc == origin:
                        route.continue_()
                    else:
                        # Shell catalog/updates are optional. Never use external
                        # services in data transformation tests.
                        route.abort()

                context.route("**/*", intercept)

            test_cases = (
                ("CSV", "synthetic-records.csv", "text/csv", CSV_FIXTURE.read_bytes()),
                ("JSON", "synthetic-records.json", "application/json",
                 json.dumps(rows, ensure_ascii=False).encode("utf-8")),
                ("XLSX", "synthetic-records.xlsx",
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", workbook_bytes(rows)),
            )
            for format_name, filename, mime, content in test_cases:
                page, errors = make_page(context, base_url)
                try:
                    data = anonymize(page, filename, mime, content)
                    if format_name == "CSV":
                        require(data.startswith(b"\xef\xbb\xbf"), "CSV output lost UTF-8 BOM")
                        actual = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"), newline="")))
                        assert_transformation(rows, actual, format_name)
                    elif format_name == "JSON":
                        actual = json.loads(data.decode("utf-8"))
                        assert_transformation(rows, actual, format_name)
                    else:
                        sheets, visibility = read_workbook(data)
                        require(list(sheets) == ["Current", "Archive"],
                                f"XLSX worksheet names/order changed: {list(sheets)}")
                        require(visibility == {"Current": "visible", "Archive": "hidden"},
                                f"XLSX worksheet visibility changed: {visibility}")
                        assert_transformation(rows, sheets["Current"], "XLSX Current")
                        assert_transformation([rows[1], rows[3]], sheets["Archive"], "XLSX Archive")
                    require(not errors, f"{format_name} browser errors: {errors}")
                    if in_memory:
                        writes = page.evaluate("window.__fixtureFetches.filter(r => r.method !== 'GET')")
                        require(not writes, f"{format_name} triggered a network write: {writes}")
                    print(f"{format_name}: PASS (downloaded bytes and unchanged fields verified)")
                finally:
                    page.close()
            require(not network_writes, f"Unexpected HTTP writes: {network_writes}")
        finally:
            browser.close()
    print(f"Marin Anonymizer browser tests: PASS ({'in-memory' if in_memory else 'HTTP'} mode)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--in-memory", action="store_true",
                        help="Render real scripts in-memory without localhost navigation")
    args = parser.parse_args()
    run(args.in_memory)
